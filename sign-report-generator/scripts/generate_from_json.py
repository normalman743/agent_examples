"""读 JSON 数据生成填好内容的立项签报 docx，大纲结构（标题/章节/子章节）由 JSON 的 sections 自由定义，不是固定模板。排版数值来自 references/format-rules-2026.json，不写死在代码里。字段说明见 README.md。

依赖 references/format-rules-2026.json（同仓库相对路径，或用数据 JSON 里的
format_config 字段指到别的文件）——这个脚本不再是单文件复制到哪都能跑，
换来的是排版标准变了只用改/换配置 JSON，不用碰这份代码。
"""

import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

_DEFAULT_RULES_PATH = Path(__file__).resolve().parent.parent / "references" / "format-rules-2026.json"


# 半角标点 -> 全角标点，只在紧挨中文字符（前或后）时转换，避免误伤纯数字/
# 英文场景（"3.5%""9:30"这类不转）。跟 check_format.py 的 check_punctuation
# 用的是同一条启发式规则，新增符号只用在这个 dict 里加一条。
_SIMPLE_PUNCT_MAP = {
    ",": "，", ";": "；", ":": "：", "!": "！", "?": "？",
}

# 引号是配对符号，不能按字符 1:1 替换（左右是不同字符），按出现顺序奇偶
# 配对：第 1、3、5...次出现转左引号，第 2、4、6...次转右引号。同一段
# 文本内配对计数，跨段落重新开始。新增配对符号（如书名号《》）在这里加
# 一条 (半角字符, 左, 右)。
_QUOTE_PAIRS = [
    ('"', "“", "”"),
    ("'", "‘", "’"),
]

_CJK_RE = re.compile(r"[一-鿿]")


def _to_fullwidth_punct(text):
    """把字符串里的半角标点转成中文全角。引号按配对顺序转左右，其余符号
    只在紧邻中文字符时转换。非字符串输入原样返回（调用方可能传数字等）。"""
    if not isinstance(text, str) or not text:
        return text

    quote_counts = {q: 0 for q, _, _ in _QUOTE_PAIRS}
    quote_map = {q: (op, cl) for q, op, cl in _QUOTE_PAIRS}

    out = []
    for i, ch in enumerate(text):
        if ch in quote_map:
            op, cl = quote_map[ch]
            quote_counts[ch] += 1
            out.append(op if quote_counts[ch] % 2 == 1 else cl)
            continue
        if ch in _SIMPLE_PUNCT_MAP:
            before = text[i - 1] if i > 0 else ""
            after = text[i + 1] if i + 1 < len(text) else ""
            if _CJK_RE.match(before) or _CJK_RE.match(after):
                out.append(_SIMPLE_PUNCT_MAP[ch])
                continue
        out.append(ch)
    return "".join(out)


def _convert_punct_any(value):
    """value 可能是字符串或数组（每项当一段），两种输入都要转换。"""
    if isinstance(value, list):
        return [_to_fullwidth_punct(v) if isinstance(v, str) else v for v in value]
    return _to_fullwidth_punct(value)


def _convert_punct_in_section(section):
    if section.get("title"):
        section["title"] = _to_fullwidth_punct(section["title"])
    if section.get("content"):
        section["content"] = _convert_punct_any(section["content"])
    table = section.get("table")
    if table:
        if table.get("headers"):
            table["headers"] = [_to_fullwidth_punct(h) if isinstance(h, str) else h for h in table["headers"]]
        if table.get("rows"):
            table["rows"] = [[_to_fullwidth_punct(c) if isinstance(c, str) else c for c in row]
                              for row in table["rows"]]
    for child in section.get("children", []):
        _convert_punct_in_section(child)


def _convert_punct_in_data(data):
    """标题/主送/正文/表格/附件/落款等所有用户可写文本字段统一转换半角
    标点为全角，渲染前做一次。attachments 名称、dept、date 也一起转，
    因为这些同样是要出现在正文里的中文文字。"""
    for key in ("title", "recipient", "intro", "dept", "date"):
        if data.get(key):
            data[key] = _convert_punct_any(data[key])
    if data.get("attachments"):
        data["attachments"] = _convert_punct_any(_attachments_list(data))
    for section in data.get("sections", []):
        _convert_punct_in_section(section)


def _load_rules(data, json_path):
    config = data.get("format_config")
    if config:
        rules_path = Path(config)
        if not rules_path.is_absolute():
            rules_path = json_path.parent / rules_path
    else:
        rules_path = _DEFAULT_RULES_PATH
    if not rules_path.exists():
        print(f"[错误] 找不到排版配置文件: {rules_path}", file=sys.stderr)
        sys.exit(1)
    with open(rules_path, encoding="utf-8") as f:
        return json.load(f)


def _font(run, name, pt_size, bold=False):
    run.font.name = name
    run.font.size = Pt(pt_size)
    run.font.bold = bold
    rp = run._r.get_or_add_rPr()
    for c in list(rp):
        if "rFonts" in c.tag:
            rp.remove(c)
    rp.append(rp.makeelement(qn("w:rFonts"), {
        qn("w:eastAsia"): name, qn("w:ascii"): name, qn("w:hAnsi"): name
    }))


def para(doc, text, font, pt=16, bold=False,
         align=None, sp_b=0, sp_a=0, line=28, indent=0, ri=0):
    """Add paragraph with Chinese document formatting.

    pt=磅值. indent=首行缩进磅值. ri=右缩进.
    line=磅值 设固定行距——全文（标题/主送/各级标题/正文/落款/附件）现在
    统一走 body.line_spacing=28。用固定行距务必让磅值大于字号，否则 Word
    "固定值"行距不跟字号撑高、会把顶部截断（28 > 标题22/正文16/表格12，安全）。
    line=None 时不设固定行距，交给 Word 按字号自动算单倍行距（保留此能力，
    目前无调用方）。
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(sp_b)
    pf.space_after = Pt(sp_a)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY   # 固定值
        pf.line_spacing = Pt(line)
    if indent:
        pf.first_line_indent = Pt(indent)
    if ri:
        pf.right_indent = Pt(ri)
    if align is not None:
        pf.alignment = align
    _font(p.add_run(text), font, pt, bold)
    return p


def _body(doc, text, rules):
    """正文内容支持两种写法：字符串（按 \\n 拆成多个自然段，一个字符串里能塞
    多段，空行 \\n\\n 会被跳过）或数组（每个元素本身就是独立一段，元素内部
    还能再用 \\n 分段）——数组写法每段是数组的一项，不用在字符串里数 \\n、
    转义引号，写起来更不容易错。两种写法混用/嵌套都可以。"""
    b = rules["body"]
    items = text if isinstance(text, list) else [text]
    for item in items:
        for line in str(item).split("\n"):
            line = line.strip()
            if line:
                para(doc, line, b["font"], b["pt"], line=b["line_spacing"], indent=b["indent"])


def _setup_page(doc, rules):
    p = rules["page"]
    s = doc.sections[0]
    s.page_width = Cm(p["width_cm"])
    s.page_height = Cm(p["height_cm"])
    s.top_margin = Cm(p["margin_top_cm"])
    s.bottom_margin = Cm(p["margin_bottom_cm"])
    s.left_margin = Cm(p["margin_left_cm"])
    s.right_margin = Cm(p["margin_right_cm"])


def _add_attachment_paragraph(doc, label, items, rules):
    """附件说明：label 与内容同一段落，多条目用真正的换行(<w:br/>)分隔，
    续行悬挂缩进到 hang。1-9 编号后面补一个空格凑对齐 10+，治标不治本——
    空格实际渲染宽度跟字体替换有关，Mac/Windows 效果可能不完全一致，
    这里先接受这个已知限制，不做更复杂的 tab stop 方案。超过 99 条直接
    报错，三位数编号这个方案完全没法凑，不硬凑一个错的对齐效果。"""
    if len(items) > 99:
        print(f"[错误] 附件有 {len(items)} 条，超过 99 条时编号对齐方式（1-9 补空格）"
              f"完全失效，需要换一种对齐实现，不能直接生成", file=sys.stderr)
        sys.exit(1)
    a = rules["attachment"]
    font, pt, indent, hang = a["font"], a["pt"], a["indent"], a["hang"]
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(rules["body"]["line_spacing"])
    pf.left_indent = Pt(hang)
    pf.first_line_indent = Pt(indent - hang)
    _font(p.add_run(label), font, pt)
    for i, item in enumerate(items):
        run = p.add_run()
        if i > 0:
            run.add_break()
        pad = " " if i + 1 < 10 else ""
        run.add_text(f"{i + 1}．{pad}{item}")
        _font(run, font, pt)
    return p


def _set_table_borders(table, size, color):
    """size 单位是 1/8 磅。"""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _shade_cell(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _add_table(doc, headers, rows, rules):
    """headers/rows 建一张单线边框表格，表头加粗居中+底纹，内容居中。"""
    t = rules["table"]
    font, pt = t["font"], t["pt"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_table_borders(table, t["border_size"], t["border_color"])

    for cell, text in zip(table.rows[0].cells, headers):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(str(text)), font, pt, bold=True)
        _shade_cell(cell, t["header_shade"])

    for row_vals in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_vals):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(str(text)), font, pt)
    return table


def _attachments_list(data):
    val = data.get("attachments")
    if not val:
        return []
    if isinstance(val, list):
        return [str(v) for v in val]
    return [str(val)]


_ATTACHMENT_REF_RE = re.compile(r"\{\{附件[:：]\s*(.+?)\s*\}\}")


def _resolve_attachment_refs(text, attachments):
    """把正文里的 {{附件:名称}} 替换成对应编号"附件N"（N 是该名称在传入的
    attachments 列表里的位置，1-based）。调用方传的是 _order_attachments 排好的
    实际引用顺序，所以编号跟着正文里第一次引用的先后走。写内容的人/agent 只用记
    附件名字、不用手数编号——增删附件或调整引用顺序后重新生成，编号自动更新。"""
    def _sub(m):
        name = m.group(1).strip()
        try:
            idx = attachments.index(name) + 1
        except ValueError:
            print(f"[警告] 正文引用了 {{{{附件:{name}}}}}，但 attachments 列表里"
                  f"找不到完全匹配的「{name}」，原样保留标记——检查是不是漏填了这条"
                  f"附件，或者名字两边打得不完全一致", file=sys.stderr)
            return m.group(0)
        return f"附件{idx}"
    return _ATTACHMENT_REF_RE.sub(_sub, str(text))


def _resolve_refs_any(value, attachments):
    """content 可能是字符串或数组（每项当一段），两种输入都要能解析
    {{附件:名称}}，所以按数组统一处理，字符串包一层再拆回来。"""
    is_list = isinstance(value, list)
    items = value if is_list else [value]
    items = [_resolve_attachment_refs(v, attachments) for v in items]
    return items if is_list else items[0]


def _resolve_refs_in_section(section, attachments):
    if section.get("content"):
        section["content"] = _resolve_refs_any(section["content"], attachments)
    for child in section.get("children", []):
        _resolve_refs_in_section(child, attachments)


def _resolve_attachment_refs_in_data(data, attachments):
    """intro/sections（含任意深度 children）里的 {{附件:名称}} 都在渲染前
    统一替换掉，attachments 本身的文字不做替换（列表项就是名称本身）。"""
    if data.get("intro"):
        data["intro"] = _resolve_refs_any(data["intro"], attachments)
    for section in data.get("sections", []):
        _resolve_refs_in_section(section, attachments)


def _collect_ref_order(data):
    """按文档顺序（intro 在前，再按 sections 深度优先）扫出 {{附件:名称}} 里每个
    名字第一次出现的先后，返回去重后的名字列表。必须在 _resolve_* 把标记替换掉
    之前调用，否则标记已经变成"附件N"、扫不到了。"""
    order, seen = [], set()

    def _scan_text(text):
        for m in _ATTACHMENT_REF_RE.finditer(str(text)):
            name = m.group(1).strip()
            if name not in seen:
                seen.add(name)
                order.append(name)

    def _scan_value(value):
        if isinstance(value, list):
            for v in value:
                _scan_text(v)
        elif value:
            _scan_text(value)

    def _scan_section(section):
        _scan_value(section.get("content"))
        for child in section.get("children", []):
            _scan_section(child)

    _scan_value(data.get("intro"))
    for section in data.get("sections", []):
        _scan_section(section)
    return order


def _order_attachments(data, attachments):
    """附件编号按正文里第一次被引用的先后（实际顺序）来，不按 attachments 数组的
    书写顺序。返回重排后的名字列表，"附件N"编号和底部附件说明都用这个顺序，保证
    正文里的"附件N"和列表第 N 条对得上。

    声明了但正文从没引用的附件仍然进列表、接在被引用的之后继续编号，但每条都会
    print 一个 warning（正常情况下每条附件都会在正文被引用一次，出现没被引用的
    多半是漏引了）。正文引用了、attachments 却没声明的名字不在这里处理，会在
    _resolve_attachment_refs 里 warn 并原样保留标记。"""
    if not attachments:
        return []
    declared = set(attachments)
    ordered = [n for n in _collect_ref_order(data) if n in declared]
    referenced = set(ordered)
    for name in attachments:
        if name not in referenced:
            print(f"[警告] 附件「{name}」在 attachments 里声明了，但正文里没有 "
                  f"{{{{附件:{name}}}}} 引用它——仍然列进附件说明并接在被引用的附件"
                  f"之后编号，确认一下是不是漏引了", file=sys.stderr)
            ordered.append(name)
    return ordered


_CN_ORDINALS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]


def _cn_number(index):
    return _CN_ORDINALS[index - 1] if index <= len(_CN_ORDINALS) else str(index)


def _add_heading(doc, text, level, index, rules):
    """四级标题字体规则对照 format-standard.md（具体数值在 format-rules-2026.json
    的 headings 里）：1级"一、"黑体不加粗，2级"（一）"楷体加粗，3级"1."仿宋加粗，
    4级"（1）"仿宋不加粗（三、四级用阿拉伯数字）。超过4级的嵌套沿用4级样式。"""
    b = rules["body"]

    h = rules["headings"].get(str(min(level, 4)), rules["headings"]["4"])
    if level == 1:
        label = f"{_cn_number(index)}、{text}"
    elif level == 2:
        label = f"（{_cn_number(index)}）{text}"
    else:
        label = f"{index}．{text}" if level == 3 else f"（{index}）{text}"
    para(doc, label, h["font"], h["pt"], bold=h["bold"], indent=b["indent"], line=b["line_spacing"])


def _render_section(doc, section, level, index, rules):
    _add_heading(doc, section.get("title", ""), level, index, rules)
    if section.get("content"):
        _body(doc, section["content"], rules)
    table = section.get("table")
    if table and table.get("headers"):
        _add_table(doc, table["headers"], table.get("rows", []), rules)
    for i, child in enumerate(section.get("children", []), 1):
        _render_section(doc, child, level + 1, i, rules)


def _warn_if_missing(data):
    if not data.get("sections"):
        print("[警告] sections 为空，正文不会有任何章节", file=sys.stderr)
    for key in ("recipient", "dept", "date"):
        if not data.get(key):
            print(f"[警告] 缺少 {key} 字段，对应内容会留空（没有默认值兜底）", file=sys.stderr)
    if data.get("closing"):
        print("[警告] 顶层 closing 字段已不支持（会被忽略，不会出现在文档里），"
              "把这段文字挪到最后一个 section 的 content 数组末尾", file=sys.stderr)
    for section in data.get("sections", []):
        _warn_if_section_has_requests(section)


def _warn_if_section_has_requests(section):
    if section.get("requests"):
        print(f"[警告] section「{section.get('title', '')}」的 requests 字段已不支持"
              f"（会被忽略，不会出现在文档里，也不会有'一是/二是'编号），改用 content 数组",
              file=sys.stderr)
    for child in section.get("children", []):
        _warn_if_section_has_requests(child)


def _title_max_chars_per_line(rules):
    """按页面可用宽度和标题字号粗估一行能放多少个全角字符，用来判断标题
    是否该手动分行——不是精确排版引擎，只是给个警戒线。"""
    page = rules["page"]
    usable_cm = page["width_cm"] - page["margin_left_cm"] - page["margin_right_cm"]
    char_width_cm = rules["title"]["pt"] * (1 / 72) * 2.54
    return int(usable_cm / char_width_cm)


def _add_title(doc, title, rules):
    """标题是数组，每个元素是独立一行。梯形/菱形怎么排、在哪断词意才完整，
    这些需要人判断语义边界，机器做不好，所以不自动算，只负责把每行单独
    居中、在明显不合理时警告（单行太长该分行了 / 分了太多行该精简标题了）。"""
    t = rules["title"]
    raw_lines = title if isinstance(title, list) else title.split("\n")
    lines = [str(line).strip() for line in raw_lines if str(line).strip()]
    if not lines:
        return

    if len(lines) == 1:
        max_chars = _title_max_chars_per_line(rules)
        if len(lines[0]) > max_chars:
            print(f"[警告] 标题单行 {len(lines[0])} 字，超过一行大约能放的 {max_chars} 字，"
                  f"建议把 title 写成数组手动分行（按梯形/菱形排布，且不能拆开词组/"
                  f"人名/地名/单位名称），或者把标题内容改短", file=sys.stderr)
    elif len(lines) > 4:
        print(f"[警告] 标题分了 {len(lines)} 行，建议不超过3-4行，太长建议精简标题内容",
              file=sys.stderr)

    for i, line in enumerate(lines):
        is_last = i == len(lines) - 1
        para(doc, line, t["font"], t["pt"], align=WD_ALIGN_PARAGRAPH.CENTER,
             sp_b=0, sp_a=(t["space_after"] if is_last else 0), line=rules["body"]["line_spacing"])


def build_document(doc, data, rules):
    _convert_punct_in_data(data)
    _warn_if_missing(data)

    attachments = _attachments_list(data)
    ordered = _order_attachments(data, attachments)
    _resolve_attachment_refs_in_data(data, ordered)

    title = data.get("title", "")
    if title:
        _add_title(doc, title, rules)

    r = rules["recipient"]
    para(doc, data.get("recipient", ""), r["font"], r["pt"],
         align=WD_ALIGN_PARAGRAPH.LEFT, sp_b=r["space_before"], sp_a=0, line=rules["body"]["line_spacing"])

    if data.get("intro"):
        _body(doc, data["intro"], rules)

    for i, section in enumerate(data.get("sections", []), 1):
        _render_section(doc, section, 1, i, rules)

    items = ordered
    if items:
        b = rules["body"]
        para(doc, "  ", b["font"], b["pt"], indent=b["indent"], line=b["line_spacing"])
        _add_attachment_paragraph(doc, "附件：", items, rules)
        # 附件说明与落款之间空两行
        para(doc, "  ", b["font"], b["pt"], indent=b["indent"], line=b["line_spacing"])
        para(doc, "  ", b["font"], b["pt"], indent=b["indent"], line=b["line_spacing"])

    f = rules["footer"]
    p = para(doc, data.get("dept", ""), f["font"], f["pt"], ri=f["right_indent"], line=rules["body"]["line_spacing"])
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = para(doc, data.get("date", ""), f["font"], f["pt"], ri=f["right_indent"], line=rules["body"]["line_spacing"])
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def main():
    if len(sys.argv) != 2:
        print("用法: python generate_from_json.py <data.json>", file=sys.stderr)
        sys.exit(1)

    json_path = Path(sys.argv[1])
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)

    rules = _load_rules(data, json_path)

    doc = Document()
    _setup_page(doc, rules)
    build_document(doc, data, rules)

    out_path = data.get("out_path")
    if out_path:
        out_path = Path(out_path)
    else:
        title = data.get("title") or "立项签报_输出"
        title = "".join(title) if isinstance(title, list) else title.replace("\n", "")
        safe_title = "".join(c for c in title if c not in '\\/:*?"<>|')
        out_path = json_path.parent / f"{safe_title}_AI草稿.docx"

    os.makedirs(out_path.parent, exist_ok=True)
    doc.save(str(out_path))
    print(f"OK  {out_path}")

    print()
    import check_format
    report, _ = check_format.check_format(out_path)
    report.print()


if __name__ == "__main__":
    main()
