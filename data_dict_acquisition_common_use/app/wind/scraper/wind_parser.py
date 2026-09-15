#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Wind 数据字典 - 解析 + 质量评分
================================

把 GUI 爬虫抓到的原始文本（wind_dict.json: {表名: 全文}）解析成结构化数据，
并给每张表打分，方便人工复查可疑表。原始文本始终保留，不破坏。

输入：
    wind/scraper/wind_dict.json        # 爬虫产出的原始文本库（{search_name: raw_text}）

输出：
    wind/dict/parsed/<表名>.json       # 每张表的结构化结果
    wind/dict/wind.json    # 全部表的结构化结果（汇总）
    wind/dict/fields.csv               # 所有表字段拉平的大表（最适合进 schema）
    wind/dict/quality_report.csv       # 每张表的状态/分数/问题，按分数升序（最该复查的在最前）

用法：
    python wind_dict_parser.py                 # 解析 scraper/wind_dict.json
    python wind_dict_parser.py <某个.json或.txt> # 额外把指定文件也纳入解析（调试用）

原始文本结构（成功加载时）：
    WDS / 万得数据服务 / [标签名拼接行]
    #<表ID>
    <中文名> - <英文名>
    简介           <- 简介标签：键值若干
    数据字典        <- 字段标签：版本/在用账号/业务主键 + 字段行(序号/中文/英文/类型/填充率/说明)
    样例数据 ...
    推荐产品 ...
    常见问题 ...
"""

import os
import sys
import csv
import json
import re

HERE = os.path.dirname(os.path.abspath(__file__))
DICT_DIR = os.path.join(HERE, "..", "dict")
RAW_JSON = os.path.join(HERE, "wind_dict.json")
SRC_XLSX = os.path.join(HERE, "exportAccountProductNew.xlsx")   # 表名->产品类型(G列状态) 来源
FIX_DIR  = os.path.join(HERE, "fix", "newjson")                # 爬虫错位表的人工修正 json

PARSED_DIR = os.path.join(DICT_DIR, "parsed")
PARSED_JSON = os.path.join(DICT_DIR, "wind.json")
FIELDS_CSV = os.path.join(DICT_DIR, "fields.csv")
REPORT_CSV = os.path.join(DICT_DIR, "quality_report.csv")
EXCEL_XLSX = os.path.join(DICT_DIR, "wind.xlsx")   # 总表：字段 / 表清单 / 质检 三个 sheet

# 识别正则
RE_SEQ = re.compile(r"^[1-9]\d*$")                     # 字段序号行：≥1（排除裸 "0"，它是 0% 填充率）
RE_ID = re.compile(r"^#(\d+)")                         # 表ID行 #101566
RE_TYPE = re.compile(r"^[A-Z][A-Z0-9_]*(\(\d+(,\d+)?\))?$")   # VARCHAR2(40) / NUMBER(20,4) / DATE
RE_FILL = re.compile(r"^(\d+(\.\d+)?%|0)$")           # 100% / 1% / 裸 0(=0%)
PERMISSION_MSG = "您尚未开通产品表权限"

# 简介标签里以全角冒号结尾的键
def _is_key(tok):
    return tok.endswith("：")


def classify(lines, raw):
    """判定加载状态：not_loaded / no_permission / ok。顺序很重要。"""
    id_line = lines[3].strip() if len(lines) > 3 else ""
    name_line = lines[4].strip() if len(lines) > 4 else ""
    m = RE_ID.match(id_line)
    table_id = m.group(1) if m else ""
    # 没加载：ID 为 0 或名字为 '-'（此时整页是空模板，常伴随权限提示）
    if table_id == "0" or name_line in ("-", ""):
        return "not_loaded", table_id, name_line
    if PERMISSION_MSG in raw:
        return "no_permission", table_id, name_line
    return "ok", table_id, name_line


def parse_intro(seg_lines):
    """简介段：按 token 流配对（键以：结尾，下一个非键 token 即其值）。"""
    toks = []
    for ln in seg_lines:
        for t in ln.split("\t"):
            t = t.strip()
            if t:
                toks.append(t)
    intro = {}
    i = 0
    while i < len(toks):
        t = toks[i]
        if _is_key(t):
            key = t.rstrip("：").strip()
            val = ""
            if i + 1 < len(toks) and not _is_key(toks[i + 1]):
                val = toks[i + 1]
                i += 1
            intro[key] = val
        i += 1
    return intro


def parse_fields(slice_lines):
    """字段段 -> [{seq,name_cn,name_en,type,fill_rate,enum_note,desc}]

    按锚点解析（不靠固定行号）：
      中文名 = 块首行；英文名 = 第二行；
      类型 = 块内第一个匹配 SQL 类型的行；填充率 = 类型之后第一个 `xx%` 行；
      类型与填充率之间的非空行 = enum_note（编码类字段的枚举类别名）；
      填充率之后的非空行 = 说明。
    """
    fields = []
    seq_pos = [k for k, l in enumerate(slice_lines) if RE_SEQ.fullmatch(l.strip())]
    seq_pos.append(len(slice_lines))  # 哨兵
    for a in range(len(seq_pos) - 1):
        start, end = seq_pos[a], seq_pos[a + 1]
        seq = int(slice_lines[start].strip())
        block = [l.strip() for l in slice_lines[start + 1:end]]
        cn = block[0] if len(block) > 0 else ""
        en = block[1] if len(block) > 1 else ""
        rest = block[2:]
        # 类型
        tpos = next((k for k, v in enumerate(rest) if RE_TYPE.match(v)), None)
        if tpos is None:
            typ, after = "", rest
        else:
            typ, after = rest[tpos], rest[tpos + 1:]
        # 填充率
        fpos = next((k for k, v in enumerate(after) if RE_FILL.match(v)), None)
        if fpos is None:
            fill, enum_note, desc = "", [x for x in after if x], []
        else:
            fill = after[fpos]
            enum_note = [x for x in after[:fpos] if x]
            desc = [x for x in after[fpos + 1:] if x]
        fields.append({
            "seq": seq,
            "name_cn": cn,
            "name_en": en,
            "type": typ,
            "fill_rate": fill,
            "enum_note": " ".join(enum_note),
            "desc": " ".join(desc),
        })
    return fields


def find_idx(lines, predicate, start=0):
    for i in range(start, len(lines)):
        if predicate(lines[i]):
            return i
    return -1


def split_sections(lines):
    """按标签把整页切成各段。返回 {标签: [该段正文行]}。顺序查找，避免误匹配。"""
    order = ["简介", "数据字典", "样例数据", "推荐产品", "常见问题"]
    idxs = {}
    pos = 5  # 跳过页头(WDS/万得数据服务/标签拼接行/#id/表名)
    for m in order:
        if m == "样例数据":
            i = find_idx(lines, lambda l: l.strip().startswith("样例数据"), pos)
        else:
            i = find_idx(lines, lambda l, mm=m: l.strip() == mm, pos)
        if i != -1:
            idxs[m] = i
            pos = i + 1
    sec = {}
    keys = list(idxs.keys())
    for j, m in enumerate(keys):
        s = idxs[m] + 1
        e = idxs[keys[j + 1]] if j + 1 < len(keys) else len(lines)
        sec[m] = lines[s:e]
    return sec


def parse_table(name, raw):
    lines = raw.splitlines()
    status, table_id, name_line = classify(lines, raw)

    result = {
        "search_name": name,
        "status": status,
        "table_id": table_id,
        "name_cn": "",
        "name_en": "",
        "version": "",
        "account": "",
        "business_key": "",
        "intro": {},
        "fields": [],
        "note": "",
        "sample_data": [],
        "recommend": [],
        "faq": [],
    }

    # 中英文表名
    if " - " in name_line:
        cn, en = name_line.split(" - ", 1)
        result["name_cn"], result["name_en"] = cn.strip(), en.strip()
    elif name_line not in ("-", ""):
        result["name_cn"] = name_line

    if status != "ok":
        return result  # 没加载/无权限：不再解析字段

    sec = split_sections(lines)

    # 简介
    result["intro"] = parse_intro(sec.get("简介", []))

    # 数据字典段：版本/在用账号/业务主键 + 字段 + 注脚
    seg = sec.get("数据字典", [])
    body_start = 0
    for k, ln in enumerate(seg):
        s = ln.strip()
        if s.startswith("版本"):
            result["version"] = (result["version"] + " " + s).strip() if result["version"] else s
            body_start = k + 1
        elif s.startswith("在用账号："):
            result["account"] = s.split("：", 1)[1].replace("与最新版本差异", "").strip()
            body_start = k + 1
        elif s.startswith("业务主键："):
            result["business_key"] = s.split("：", 1)[1].strip()
            body_start = k + 1
        elif RE_SEQ.fullmatch(s):
            break
    end = len(seg)
    for k in range(body_start, len(seg)):
        if seg[k].strip().startswith("注："):
            end = k
            result["note"] = seg[k].strip()
            break
    result["fields"] = parse_fields(seg[body_start:end])

    # 样例数据 / 推荐产品 / 常见问题（原样保留非空行，过滤"暂无数据"）
    def clean(seg_lines):
        out = [s.strip() for s in seg_lines if s.strip()]
        return [] if out == ["暂无数据"] else out

    result["sample_data"] = clean(sec.get("样例数据", []))
    result["recommend"] = clean(sec.get("推荐产品", []))
    result["faq"] = clean(sec.get("常见问题", []))

    return result


def score_table(t):
    """完整度评分(0-100) + 问题清单。数据字典段占 55 分(>50%)。分数越低越该复查。

    数据字典段 55： 字段结构完整度 40 + 序号连续 8 + 注脚完整 7
    元信息     15： 表名中英 5 + 表ID 3 + 业务主键 4 + 版本 2 + 账号 1
    简介       10： 键值齐全度
    样例数据   10
    推荐+常见问题 10： 各 5
    """
    issues = []
    status = t["status"]
    if status == "not_loaded":
        return 0, ["未加载(#0/-)，需重抓"], True
    if status == "no_permission":
        return 30, ["无权限，需申请字典查阅"], True

    score = 0.0
    # ---- 数据字典字段段 55 ----
    fields = t["fields"]
    if not fields:
        issues.append("字段数为0")
    else:
        no_cn = no_en = bad_type = 0
        acc = 0.0
        for f in fields:
            parts = 0
            if f["name_cn"]:
                parts += 1
            else:
                no_cn += 1
            if f["name_en"]:
                parts += 1
            else:
                no_en += 1
            if f["type"] and RE_TYPE.match(f["type"]):
                parts += 1
            else:
                bad_type += 1
            acc += parts / 3
        score += 40 * (acc / len(fields))
        seqs = [f["seq"] for f in fields]
        if seqs == list(range(1, len(seqs) + 1)):
            score += 8
        else:
            issues.append(f"序号不连续: {seqs}")
        if t["note"]:
            score += 7
        else:
            issues.append("缺注脚(可能复制截断)")
        if no_cn:
            issues.append(f"{no_cn}个字段缺中文名")
        if no_en:
            issues.append(f"{no_en}个字段缺英文名")
        if bad_type:
            issues.append(f"{bad_type}个字段类型异常")

    # ---- 元信息 15 ----
    if t["name_cn"] and t["name_en"]:
        score += 5
    else:
        issues.append("表名中/英文未完整解析")
    if t["table_id"]:
        score += 3
    else:
        issues.append("缺表ID")
    if t["business_key"]:
        score += 4
    else:
        issues.append("缺业务主键")
    if t["version"]:
        score += 2
    if t["account"]:
        score += 1

    # ---- 简介 10 ----
    if t["intro"]:
        score += min(10, len(t["intro"]))
    else:
        issues.append("简介为空")
    # ---- 样例数据 10 ----
    if t["sample_data"]:
        score += 10
    else:
        issues.append("无样例数据")
    # ---- 推荐+常见问题 10 ----
    if t["recommend"]:
        score += 5
    if t["faq"]:
        score += 5

    score = round(min(100, score))
    # 只对"数据字典/元信息"层面的问题标记需复查；缺样例/简介等可选项不强制复查
    review_kw = ("字段", "序号", "注脚", "表名", "表ID", "业务主键")
    needs_review = any(any(k in i for k in review_kw) for i in issues)
    return score, issues, needs_review


def load_product_type():
    """从 exportAccountProductNew.xlsx 读取 {表名(B列) -> 产品类型(G列状态)}。"""
    if not os.path.exists(SRC_XLSX):
        print(f"[!] 未找到 {SRC_XLSX}，产品类型列将为空")
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(SRC_XLSX, read_only=True, data_only=True)
    ws = wb.active
    pmap = {}
    for r in ws.iter_rows(min_row=2, min_col=2, max_col=7, values_only=True):
        name, ptype = r[0], r[5]   # B=表名, G=状态
        if name and str(name).strip():
            pmap[str(name).strip()] = "" if ptype is None else str(ptype).strip()
    return pmap


def load_raw_store():
    if not os.path.exists(RAW_JSON):
        sys.exit(f"[-] 找不到原始库：{RAW_JSON}（先跑爬虫）")
    with open(RAW_JSON, "r", encoding="utf-8") as f:
        return json.load(f)


def main():
    store = load_raw_store()

    # 允许额外纳入调试文件（.txt 单表 / .json 多表）
    for extra in sys.argv[1:]:
        if extra.endswith(".txt") and os.path.exists(extra):
            key = os.path.splitext(os.path.basename(extra))[0]
            store[key] = open(extra, encoding="utf-8").read()
        elif extra.endswith(".json") and os.path.exists(extra):
            store.update(json.load(open(extra, encoding="utf-8")))

    os.makedirs(PARSED_DIR, exist_ok=True)
    pmap = load_product_type()      # 表名 -> 产品类型(G列状态)
    parsed_all = {}
    report_rows = []
    field_rows = []

    mismatch_fixed   = []   # search_name != name_en，有 fix
    mismatch_no_fix  = []   # search_name != name_en，无 fix

    for name, raw in store.items():
        t = parse_table(name, raw)

        # 爬虫点错结果导致 search_name != name_en：尝试用人工修正文件覆盖
        if t.get("name_en", "") and t["name_en"] != name:
            fix_path = os.path.join(FIX_DIR, f"{name}.json")
            if os.path.exists(fix_path):
                with open(fix_path, encoding="utf-8") as fh:
                    t = json.load(fh)
                t["search_name"] = name          # 保留原始 search_name
                mismatch_fixed.append((name, t["name_en"]))
            else:
                mismatch_no_fix.append((name, t["name_en"]))
                t.setdefault("issues", []).append(
                    f"search_name({name}) != name_en({t['name_en']})，无修正文件，字段可能错位"
                )
                t["needs_review"] = True

        score, issues, needs_review = score_table(t)
        t["score"] = score
        t["issues"] = issues
        t["needs_review"] = needs_review
        t["product_type"] = pmap.get(name, "")   # G列状态：正式/试用/测试/样例
        parsed_all[name] = t

        safe = "".join(c if (c.isalnum() or c in " _-（）()") else "_" for c in name).strip()[:120] or name
        with open(os.path.join(PARSED_DIR, f"{safe}.json"), "w", encoding="utf-8") as f:
            json.dump(t, f, ensure_ascii=False, indent=2)

        report_rows.append({
            "search_name": name,
            "status": t["status"],
            "score": score,
            "needs_review": "是" if needs_review else "",
            "product_type": t["product_type"],
            "table_id": t["table_id"],
            "name_cn": t["name_cn"],
            "name_en": t["name_en"],
            "field_count": len(t["fields"]),
            "issues": "; ".join(issues),
        })

        for f in t["fields"]:
            field_rows.append({
                "search_name": name,
                "product_type": t["product_type"],
                "table_id": t["table_id"],
                "table_cn": t["name_cn"],
                "table_en": t["name_en"],
                "seq": f["seq"],
                "name_cn": f["name_cn"],
                "name_en": f["name_en"],
                "type": f["type"],
                "fill_rate": f["fill_rate"],
                "enum_note": f.get("enum_note", ""),
                "desc": f["desc"],
            })

    def canonicalize_wind(tables_dict):
        def _join_path(*segs):
            parts = [str(s).strip() for s in segs if s and str(s).strip()]
            return " / ".join(parts)
            
        canonical_tables = []
        for key, r in tables_dict.items():
            intro = r.get("intro", {}) or {}
            fields = []
            for c in r.get("fields", []) or []:
                fields.append({
                    "seq": c.get("seq", 0),
                    "col": c.get("name_en", ""),
                    "name": c.get("name_cn", ""),
                    "type": c.get("type", ""),
                    "unit": "",
                    "enum_ref": c.get("enum_note", ""),
                    "status": "",
                    "fill_rate": c.get("fill_rate", ""),
                    "desc": c.get("desc", ""),
                })
            base = {
                "table_cn": r.get("name_cn", ""),
                "table_en": r.get("name_en", key),
                "primary_key": r.get("business_key", ""),
                "biz_key": r.get("business_key", ""),
                "status": r.get("status", ""),
                "update_time": "",
                "update_freq": "",
                "desc": intro.get("产品说明", ""),
                "category": intro.get("所属模块", ""),
                "path": _join_path(intro.get("所属数据库"), intro.get("所属模块")),
                "fields": fields,
                "enums": {},
                "model_id": None,
                "obj_id": r.get("table_id")
            }
            for k in ("intro", "sample_data", "note", "recommend", "faq", "score", "issues", "product_type", "version", "account", "needs_review", "search_name"):
                if k in r:
                    base[k] = r[k]
            canonical_tables.append(base)
        return canonical_tables

    out_data = {
        "database": "万得（Wind）数据字典",
        "database_en": "wind",
        "source": "万得金融终端",
        "updated_at": "2026-06-18",
        "count": len(parsed_all),
        "tables": canonicalize_wind(parsed_all)
    }
    with open(PARSED_JSON, "w", encoding="utf-8") as f:
        json.dump(out_data, f, ensure_ascii=False, indent=2)

    with open(FIELDS_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["search_name", "product_type", "table_id", "table_cn", "table_en",
                                          "seq", "name_cn", "name_en", "type", "fill_rate", "enum_note", "desc"])
        w.writeheader(); w.writerows(field_rows)

    report_rows.sort(key=lambda r: r["score"])  # 低分在前
    with open(REPORT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["search_name", "status", "score", "needs_review", "product_type",
                                          "table_id", "name_cn", "name_en", "field_count", "issues"])
        w.writeheader(); w.writerows(report_rows)

    # 总 Excel：字段 / 表清单 / 质检 三个 sheet
    write_excel(parsed_all, field_rows, report_rows)

    # 控制台小结
    from collections import Counter
    c = Counter(t["status"] for t in parsed_all.values())
    nrev = sum(1 for t in parsed_all.values() if t["needs_review"])
    print(f"[✓] 解析 {len(parsed_all)} 张表  状态: {dict(c)}")
    print(f"    需人工复查: {nrev} 张  ->  {REPORT_CSV}")
    print(f"    字段总数: {len(field_rows)}  ->  {FIELDS_CSV}")
    print(f"    结构化汇总(总JSON) -> {PARSED_JSON}")
    print(f"    总 Excel -> {EXCEL_XLSX}")

    if mismatch_fixed:
        print(f"\n[!] search_name != name_en，已用 fix 覆盖（共 {len(mismatch_fixed)} 张）：")
        for sn, en in mismatch_fixed:
            print(f"      {sn}  ->  {en}")
    if mismatch_no_fix:
        print(f"\n[!] search_name != name_en，无修正文件，字段可能错位（共 {len(mismatch_no_fix)} 张）：")
        for sn, en in mismatch_no_fix:
            print(f"      {sn}  ->  {en}  ← 需补 fix/newjson/{sn}.json")


def write_excel(parsed_all, field_rows, report_rows):
    """把全部结果写进一个 .xlsx：字段总表 / 表清单 / 质检报告 三个 sheet。"""
    import openpyxl

    wb = openpyxl.Workbook()

    # sheet1: 字段总表（所有表字段拉平）
    ws = wb.active
    ws.title = "字段总表"
    cols = ["search_name", "product_type", "table_id", "table_cn", "table_en",
            "seq", "name_cn", "name_en", "type", "fill_rate", "enum_note", "desc"]
    ws.append(cols)
    for r in field_rows:
        ws.append([r.get(c, "") for c in cols])

    # sheet2: 表清单（每张表一行的元信息汇总）
    ws2 = wb.create_sheet("表清单")
    cols2 = ["search_name", "product_type", "status", "score", "needs_review", "table_id",
             "name_cn", "name_en", "version", "account", "business_key",
             "field_count", "sample_rows", "issues"]
    ws2.append(cols2)
    for name, t in sorted(parsed_all.items(), key=lambda kv: kv[1]["score"]):
        ws2.append([
            name, t.get("product_type", ""), t["status"], t["score"], "是" if t["needs_review"] else "",
            t["table_id"], t["name_cn"], t["name_en"], t["version"],
            t["account"], t["business_key"], len(t["fields"]),
            len(t["sample_data"]), "; ".join(t["issues"]),
        ])

    # sheet3: 质检报告（按分数升序，最该复查的在前）
    ws3 = wb.create_sheet("质检报告")
    cols3 = ["search_name", "status", "score", "needs_review", "product_type",
             "table_id", "name_cn", "name_en", "field_count", "issues"]
    ws3.append(cols3)
    for r in report_rows:
        ws3.append([r.get(c, "") for c in cols3])

    wb.save(EXCEL_XLSX)


if __name__ == "__main__":
    main()
